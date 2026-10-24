#!/usr/bin/env python3
"""
postprocess_docx_tables.py

Apply consistent styling to all tables in a DOCX file:
  - Uniform borders (color + width)
  - Header row: background color, white bold text
  - Body rows: alternating background colors
  - All cells: padding, vertical centering, horizontal centering
  - Preserve inline-code run formatting inside table cells

Usage:
    python postprocess_docx_tables.py output.docx
    python postprocess_docx_tables.py output.docx \
        --header-bg 00174F --header-fg FFFFFF \
        --row-alt EDFDFD --border FFFFFF \
        --border-width 2.25 \
        --font "Liberation Sans" --body-size 10.5 --text-color 2C3E50
"""

import argparse
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


def _hex(c):
    return c.lstrip('#').upper()


def _remove_children(parent, tag):
    for child in parent.findall(qn(tag)):
        parent.remove(child)


def set_table_borders(table, color_hex, width_pt):
    """
    Set uniform borders on the whole table. w:sz is in eighths of a point,
    valid range 2..96 for 'single' style borders.
    """
    sz = int(round(width_pt * 8))
    sz = max(2, min(96, sz))
    tblPr = table._tbl.tblPr
    _remove_children(tblPr, 'w:tblBorders')
    tblPr.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top      w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:left     w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:bottom   w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:right    w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:insideH  w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'<w:insideV  w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
        f'</w:tblBorders>'
    ))


def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    _remove_children(tcPr, 'w:shd')
    tcPr.append(parse_xml(
        f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{color_hex}"/>'
    ))


def set_cell_vertical_align(cell, val="center"):
    tcPr = cell._tc.get_or_add_tcPr()
    _remove_children(tcPr, 'w:vAlign')
    tcPr.append(parse_xml(
        f'<w:vAlign {nsdecls("w")} w:val="{val}"/>'
    ))


def set_cell_margins(cell, top=40, left=80, bottom=40, right=80):
    """Cell padding in twips (1/20 pt). 80 twips = 4pt; 40 twips = 2pt."""
    tcPr = cell._tc.get_or_add_tcPr()
    _remove_children(tcPr, 'w:tcMar')
    tcPr.append(parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top    w:w="{top}"    w:type="dxa"/>'
        f'<w:left   w:w="{left}"   w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:right  w:w="{right}"  w:type="dxa"/>'
        f'</w:tcMar>'
    ))


def is_inline_code_run(run):
    """Return True for runs marked with a Word character style for inline code.

    Pandoc normally emits inline code with the ``VerbatimChar`` character
    style. Also recognize common custom style names so custom reference
    documents and Lua filters can use names such as ``InlineCode`` or ``Code``.
    """
    r_pr = run._r.rPr
    if r_pr is None:
        return False

    r_style = r_pr.find(qn('w:rStyle'))
    if r_style is None:
        return False

    style_id = r_style.get(qn('w:val'), '')
    normalized = ''.join(ch for ch in style_id.lower() if ch.isalnum())
    return (
        'verbatim' in normalized
        or 'inlinecode' in normalized
        or normalized in {'code', 'codechar', 'sourcecode', 'sourcecodechar'}
        or normalized.endswith('codechar')
    )


def style_runs(paragraph, color_hex, bold, font_name, font_size_pt):
    for run in paragraph.runs:
        # Inline-code runs have their own typography, foreground color, and
        # often a contrasting background. Do not overwrite those run-level
        # properties with table body/header styling.
        if is_inline_code_run(run):
            continue

        run.font.color.rgb = RGBColor.from_string(color_hex)
        run.font.bold = bold
        run.font.name = font_name
        run.font.size = Pt(font_size_pt)

        # Ensure the font is applied to all script ranges, not just Latin
        rPr = run._r.get_or_add_rPr()
        _remove_children(rPr, 'w:rFonts')
        rPr.insert(0, parse_xml(
            f'<w:rFonts {nsdecls("w")} '
            f'w:ascii="{font_name}" w:hAnsi="{font_name}" '
            f'w:cs="{font_name}" w:eastAsia="{font_name}"/>'
        ))


def is_header_row(row):
    """Pandoc marks header rows with <w:trPr><w:tblHeader/></w:trPr>."""
    trPr = row._tr.find(qn('w:trPr'))
    if trPr is not None and trPr.find(qn('w:tblHeader')) is not None:
        return True
    return False


def postprocess_docx_tables(docx_path, opts):
    doc = Document(docx_path)

    tables_processed = 0
    cells_styled = 0

    for table in doc.tables:
        tables_processed += 1
        set_table_borders(table, _hex(opts.border), opts.border_width)

        body_row_index = 0
        for row in table.rows:
            if is_header_row(row):
                bg = _hex(opts.header_bg)
                text_color = _hex(opts.header_fg)
                bold = True
            else:
                body_row_index += 1
                bg = 'FFFFFF' if (body_row_index % 2 == 1) else _hex(opts.row_alt)
                text_color = _hex(opts.text_color)
                bold = False

            # Merged cells: python-docx returns the same tc for each grid
            # position it spans, so dedupe to avoid wasted work.
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))

                set_cell_shading(cell, bg)
                set_cell_vertical_align(cell, 'center')
                set_cell_margins(cell)

                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    style_runs(paragraph, text_color, bold,
                               opts.font, opts.body_size)
                cells_styled += 1

    doc.save(docx_path)
    print(f"  ✓ Processed {tables_processed} table(s), styled {cells_styled} cell(s)")
    print(f"  ✓ Wrote {docx_path}")


def main():
    parser = argparse.ArgumentParser(
        description="Apply custom styling to tables in a DOCX file."
    )
    parser.add_argument("docx_file")

    parser.add_argument("--header-bg",    default="00174F")
    parser.add_argument("--header-fg",    default="FFFFFF")
    parser.add_argument("--row-alt",      default="EDFDFD")
    parser.add_argument("--border",       default="FFFFFF")
    parser.add_argument("--text-color",   default="2C3E50")
    parser.add_argument("--font",         default="Liberation Sans")
    parser.add_argument("--body-size",    type=float, default=10.5)
    parser.add_argument("--border-width", type=float, default=2.25)

    args = parser.parse_args()

    try:
        postprocess_docx_tables(args.docx_file, args)
    except FileNotFoundError:
        print(f"Error: file not found: {args.docx_file}", file=sys.stderr)
        sys.exit(1)
    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()